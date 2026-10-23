"""Decode headers and resource identities only; no molecular numeric values."""
import collections
import hashlib
import json
from pathlib import Path
import openpyxl

OUT=Path(__file__).resolve().parents[2]/'log/20261022/decrypte_feedback_pilot/molecular'
listing=json.loads((OUT/'FILELIST.json').read_text())['row_data']
jurkat=[r for r in listing if r['relative_path'].startswith('other/Dose response data - Jurkat proteome/') and r['name'].endswith('_ALL.txt')]
search=[r for r in listing if r['relative_path'].startswith('search/Search results - Jurkat proteome')]
book=openpyxl.load_workbook((OUT/'mapping.body').open('rb'),read_only=True,data_only=False)
mapping=[]
for sheet in book:
    rows=[]
    for row in sheet.iter_rows(min_row=1,max_row=3):
        rows.append([{'column':c.column,'type':c.data_type,'text':c.value if isinstance(c.value,str) else '[nontext redacted]'} for c in row if c.value is not None])
    mapping.append({'sheet':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'text_first3rows':rows})
raw=(OUT/'jurkat_jq1.body').read_bytes()
header=raw.split(b'\n',1)[0].decode('utf-8-sig').rstrip('\r').split('\t')
result={'scope':'metadata and first header line only; molecular value rows never decoded',
        'jurkat_all_count':len(jurkat),'jurkat_all_total_bytes':sum(r['size'] for r in jurkat),
        'jurkat_all_drugs':[r['relative_path'].split('/')[2] for r in jurkat],
        'downloaded_jq1_header':header,'downloaded_jq1_newline_count':raw.count(b'\n'),
        'downloaded_jq1_carriage_count':raw.count(b'\r'),
        'search_results':search,'mapping_headers':mapping}
(OUT/'MOLECULAR_HEADER_CONTRACT.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps({'jurkat_all_count':len(jurkat),'jurkat_all_total_bytes':sum(r['size'] for r in jurkat),
                  'newlines':raw.count(b'\n'),'carriage':raw.count(b'\r'),
                  'search_files':[(r['name'],r['size']) for r in search],'mapping':mapping},indent=2))
