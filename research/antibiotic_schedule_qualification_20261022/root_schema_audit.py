"""Two pinned Git blobs; export schema and counts, never response values."""
import base64
import csv
import hashlib
import io
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
import re

import openpyxl
import requests

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "log/20261022/antibiotic_schedule_qualification/schema_audit"
FILES = [
    ("growth_data.csv", "4197e4588383cd2a2441704f758f2533a3abe6bf", 17137),
    ("multiplexing_control.xlsx", "0dc02c1c85a7d10d6783a361d2ae41f91d9bb685", 2029932),
]
HEADER = re.compile(r"(?:time|hour|hours|treatment|condition|control|growth|drug|dose|concentration|chamber|valve|input|output|sequence|combinatorial input|combination|A|B|A\+B|A-B|B-A|amikacin|cefepime|clindamycin|levofloxacin|nitrofurantoin|rifampicin|sulfamethoxazole|tetracycline|trimethoprim)(?:[ _\-a-zA-Z()%/+]*)", re.I)


def headers(row):
    return [str(x).strip() if isinstance(x, str) and HEADER.fullmatch(x.strip())
            and not any(c.isdigit() for c in x) else
            None if x is None or x == "" else "<redacted>" for x in row]


def numeric(value):
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


def csv_schema(payload):
    rows = list(csv.reader(io.StringIO(payload.decode("utf-8-sig"))))
    return {"rows_including_header": len(rows),
            "row_width_counts": dict(Counter(len(row) for row in rows)),
            "first_row_header_projection": headers(rows[0]) if rows else [],
            "body_numeric_cells": sum(numeric(x) for row in rows[1:] for x in row),
            "body_blank_cells": sum(x == "" for row in rows[1:] for x in row)}


def workbook_schema(payload):
    book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    result = []
    for sheet in book:
        counts = Counter()
        first = []
        for index, row in enumerate(sheet.iter_rows()):
            if index == 0:
                first = headers([cell.value for cell in row])
            for cell in row:
                counts["blank" if cell.value is None else cell.data_type] += 1
        result.append({"sheet_name": sheet.title, "rows": sheet.max_row,
                       "columns": sheet.max_column, "first_row_header_projection": first,
                       "cell_type_counts": dict(counts)})
    book.close()
    return result


def checks():
    assert headers(["1.25", 1.25, "unrecognised", "Time", "A+B"]) == ["<redacted>", "<redacted>", "<redacted>", "Time", "A+B"]
    assert headers(["Control", "Control", None]) == ["Control", "Control", None]
    projected = csv_schema(b"Time,Control\n0,9.87654321\n1,8.76543210\n")
    assert "9.87654321" not in json.dumps(projected) and projected["rows_including_header"] == 3
    return {"distinct_export_checks_passed": 3}


def main():
    if (OUT / "NETWORK.json").exists():
        raise FileExistsError("No duplicate acquisition or receipt overwrite")
    (OUT / "sealed_sources").mkdir(parents=True, exist_ok=True)
    receipts, results, consumed = [], [], 0
    (OUT / "EXPORT_CHECKS.json").write_text(json.dumps(checks(), indent=2), encoding="utf-8")
    for name, sha, size in FILES:
        url = f"https://api.github.com/repos/Uchicago-TAY/LIST/git/blobs/{sha}"
        receipt = {"file": name, "url": url, "utc": datetime.now(timezone.utc).isoformat()}
        body = b""
        try:
            with requests.get(url, timeout=45, stream=True, allow_redirects=False) as response:
                body = response.raw.read(3_000_000 - consumed, decode_content=True)
                receipt["status"] = response.status_code
            consumed += len(body)
            receipt.update(bytes=len(body), sha256=hashlib.sha256(body).hexdigest())
            (OUT / "sealed_sources" / (name + ".http.json")).write_bytes(body)
            if receipt["status"] != 200 or consumed >= 3_000_000:
                raise ValueError("HTTP failure or acquisition ceiling; do not parse")
            obj = json.loads(body)
            payload = base64.b64decode(obj["content"], validate=False)
            actual = hashlib.sha1(b"blob " + str(len(payload)).encode() + b"\0" + payload).hexdigest()
            if actual != sha or len(payload) != size or obj["sha"] != sha:
                raise ValueError("Pinned source mismatch")
            receipt.update(blob_sha_verified=True, decoded_blob_bytes=len(payload))
            result = csv_schema(payload) if name.endswith(".csv") else workbook_schema(payload)
            results.append({"file": name, "git_sha": sha, "schema": result})
        except Exception as exc:
            receipt["error_type"] = type(exc).__name__
            receipt["error"] = str(exc)
        receipts.append(receipt)
        (OUT / "NETWORK.json").write_text(json.dumps(receipts, indent=2), encoding="utf-8")
    exported = {"files": results, "requests": len(receipts), "http_body_bytes": consumed,
                "response_values_exported": False,
                "code_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest()}
    (OUT / "SCHEMA.json").write_text(json.dumps(exported, indent=2), encoding="utf-8")
    print(json.dumps(exported, indent=2))


if __name__ == "__main__":
    main()
