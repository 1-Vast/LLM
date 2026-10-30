"""Optional next-step cleaner: build a partial structural map from public AZ/GDSC data.

Writes one row per matrix record, with:
  cell/drug/dose/plate/cmatrix/QC metadata and candidate QC-date groups,
  and explicit caveats that QC_UPLOAD_DATE is not a treatment date and the
  candidate "occasion"/"technical replicate" indices are proxies only.

This is NOT an effect analysis and does NOT certify independent cultures.
It is provided as the next cleaning/reconciliation script; it was not executed in
this task (the task required no effect execution; the structural reconstruction is
already verified in data3_verify.json).
"""
import csv
import io
import os
import unicodedata
import zipfile
from collections import defaultdict

import openpyxl
import pandas as pd

B = r'D:\MAESTRO'
TMP = os.path.join(B, 'tmp')
RAW = os.path.join(TMP, 'extq1', 'gsdc_well_stats.zip')
SUPP = os.path.join(TMP, 'epmc_supp_list.txt')
OUT = os.path.join(B, r'research\2026-10-28-az-download-verify\partial_map.csv')


def norm(s):
    s = unicodedata.normalize('NFKC', str(s)).casefold()
    return ''.join(ch for ch in s if ch.isalnum())


def as_id(x):
    s = str(x or '').strip()
    return s[:-2] if s.endswith('.0') else s


id2name = {}


def add(did, name):
    did = as_id(did); name = str(name or '').strip()
    if did and name and did not in id2name:
        id2name[did] = name


with zipfile.ZipFile(SUPP) as z:
    wb = openpyxl.load_workbook(z.open('EMS194010-supplement-Table_S2.xlsx'), read_only=True, data_only=True)
    for sh in wb.sheetnames:
        ws = wb[sh]; it = ws.iter_rows(values_only=True); hdr = next(it)
        ix = {str(c).strip(): i for i, c in enumerate(hdr)}
        for row in it:
            def cv(c):
                i = ix.get(c); return row[i] if i is not None and i < len(row) else None
            if 'Compound1_ID' in ix:
                add(cv('Compound1_ID'), cv('Compound1_name')); add(cv('Compound2_ID'), cv('Compound2_name'))
            if 'Drug_ID' in ix:
                add(cv('Drug_ID'), cv('Drug_name'))
    wb.close()
for f in ['gdsc-007_matrix_results.csv', 'gdsc-008_matrix_results.csv', 'gdsc-009_matrix_results.csv',
          'gdsc010_matrix_results.csv', 'sandpiper-01_matrix_results.csv']:
    p = os.path.join(TMP, 'extq1', f)
    d = pd.read_csv(p, usecols=['lib1_ID', 'lib1_name', 'lib2_ID', 'lib2_name'], low_memory=False)
    for a, an, b, bn in d.itertuples(index=False):
        add(a, an); add(b, bn)

qc = {}
for fn in ['plate_qc_Sandpiper-01_13May22.csv', 'plate_qc_Sandpiper-02_13May22.csv']:
    d = pd.read_csv(os.path.join(TMP, 'extq1', fn), usecols=['BARCODE', 'SCAN_ID', 'QC_UPLOAD_DATE', 'QC_FLAG'])
    for r in d.itertuples(index=False):
        qc[as_id(r.BARCODE)] = (as_id(r.SCAN_ID), str(r.QC_UPLOAD_DATE), str(r.QC_FLAG))

# pass 1: group matrix records per (cell, pair)
pair_records = defaultdict(list)
headers = None
with zipfile.ZipFile(RAW) as z:
    for mn in [i.filename for i in z.infolist() if i.filename.lower().endswith('.csv')]:
        with z.open(mn) as f:
            txt = io.TextIOWrapper(f, encoding='utf-8', errors='replace', newline='')
            rdr = csv.reader(txt); hdr = next(rdr)
            if headers is None:
                headers = hdr
            ix = {c: i for i, c in enumerate(hdr)}
            for row in rdr:
                cell = norm(row[ix['CELL_LINE_NAME']])
                d1 = norm(id2name.get(as_id(row[ix['lib1_drug_id']]), ''))
                d2 = norm(id2name.get(as_id(row[ix['lib2_drug_id']]), ''))
                if not d1 or not d2:
                    continue
                pair = tuple(sorted((d1, d2)))
                bc = as_id(row[ix['BARCODE']]); cm = as_id(row[ix['cmatrix']])
                pair_records[(cell, pair)].append({
                    'cell_line_name': row[ix['CELL_LINE_NAME']],
                    'barcode': bc, 'cmatrix': cm, 'drugset_id': as_id(row[ix['DRUGSET_ID']]),
                    'project_name': row[ix['PROJECT_NAME']],
                    'lib1_drug_id': as_id(row[ix['lib1_drug_id']]), 'lib1_name': id2name.get(as_id(row[ix['lib1_drug_id']]), ''),
                    'lib1_dose': row[ix['lib1_dose']], 'lib1_conc': row[ix['lib1_conc']],
                    'lib2_drug_id': as_id(row[ix['lib2_drug_id']]), 'lib2_name': id2name.get(as_id(row[ix['lib2_drug_id']]), ''),
                    'lib2_dose': row[ix['lib2_dose']], 'lib2_conc': row[ix['lib2_conc']],
                    'scan_id': qc.get(bc, ('', '', ''))[0], 'qc_upload_date': qc.get(bc, ('', '', ''))[1],
                    'qc_flag': qc.get(bc, ('', '', ''))[2],
                })

# pass 2: write per-matrix rows with candidate QC-date groups (proxies only)
cols = ['cell_line_name', 'barcode', 'scan_id', 'qc_upload_date', 'qc_flag', 'project_name', 'drugset_id', 'cmatrix',
        'lib1_drug_id', 'lib1_name', 'lib1_dose', 'lib1_conc', 'lib2_drug_id', 'lib2_name', 'lib2_dose', 'lib2_conc',
        'n_matrix_records_for_pair', 'qc_date_occasion_proxy', 'technical_replicate_proxy', 'mapping_status', 'certified_independent_culture']
with open(OUT, 'w', newline='', encoding='utf-8') as f:
    w = csv.DictWriter(f, fieldnames=cols)
    w.writeheader()
    for (cell, pair), recs in pair_records.items():
        # unique matrix records (barcode,cmatrix), sorted by QC date/barcode/cmatrix
        uniq = {}
        for r in recs:
            uniq.setdefault((r['barcode'], r['cmatrix']), r)
        by_date = defaultdict(list)
        for (bc, cm), r in uniq.items():
            by_date[r['qc_upload_date']].append(r)
        date_ids = {d: i + 1 for i, d in enumerate(sorted(by_date))}
        exact = len(by_date) == 6 and all(len(v) == 3 for v in by_date.values())
        for d in sorted(by_date):
            for j, r in enumerate(sorted(by_date[d], key=lambda x: (x['barcode'], x['cmatrix']))):
                r = dict(r)
                r['n_matrix_records_for_pair'] = len(uniq)
                r['qc_date_occasion_proxy'] = date_ids[d]
                r['technical_replicate_proxy'] = j + 1
                r['mapping_status'] = 'exact_6_qc_dates_x3' if exact else ('six_qc_dates_not3' if len(by_date) == 6 else 'not_6x3')
                r['certified_independent_culture'] = 'FALSE'
                w.writerow({k: r.get(k, '') for k in cols})
print('wrote', OUT, 'pair units', len(pair_records))
