"""R204-DATA3 local verification and mapping reconstruction (no effects).

Streams the public raw well data, rebuilds canonical (cell,pair,cmatrix) matrix records,
recomputes the X2 structural counts, attempts 6-QC-date x 3-matrix reconstruction,
checks field availability and public-code compatibility, and gives structural readiness
statuses.  Does NOT read/analyse response effects; no fits/statistics on responses.
Writes data3_verify.json (sorted keys, no timestamps).
"""
import csv
import hashlib
import io
import json
import os
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict

import pandas as pd
import openpyxl

B = r'D:\MAESTRO'
HERE = os.path.join(B, r'research\2026-10-28-az-download-verify')
TMP = os.path.join(B, 'tmp')
MANIFEST = os.path.join(HERE, 'data3_manifest.json')
PROTOCOL = os.path.join(HERE, 'PROTOCOL_DATA3.md')
OUT = os.path.join(HERE, 'data3_verify.json')


def sha(p):
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        for b in iter(lambda: f.read(1 << 22), b''):
            h.update(b)
    return h.hexdigest().upper()


def norm(s):
    s = unicodedata.normalize('NFKC', str(s)).casefold()
    return ''.join(ch for ch in s if ch.isalnum())


def as_id(x):
    if x is None:
        return ''
    s = str(x).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s


def q(x):
    if isinstance(x, dict):
        return {k: q(v) for k, v in x.items()}
    if isinstance(x, (list, tuple)):
        return [q(v) for v in x]
    if isinstance(x, float):
        return round(x, 6) if x == x else None
    return x


man = json.load(open(MANIFEST, encoding='utf-8'))
entries = {e['logical']: e for e in man['entries']}
raw_path = os.path.join(B, entries['figshare_raw_well_zip']['path_rel'])
supp_path = os.path.join(B, entries['europepmc_supplementary_zip']['path_rel'])

# canonical drug ID -> name
id2name = {}
name_conflicts = Counter()
sources = Counter()


def add_map(did, name, src):
    did = as_id(did); name = str(name).strip() if name is not None else ''
    if not did or not name or name.lower() == 'nan':
        return
    if did in id2name and norm(id2name[did]) != norm(name):
        name_conflicts[did] += 1
        return
    id2name[did] = name
    sources[src] += 1


with zipfile.ZipFile(supp_path) as z:
    wb = openpyxl.load_workbook(z.open('EMS194010-supplement-Table_S2.xlsx'), read_only=True, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        hdr = next(rows)
        idx = {str(c).strip(): i for i, c in enumerate(hdr)}
        for row in rows:
            def cv(col):
                i = idx.get(col)
                return row[i] if i is not None and i < len(row) else None
            if 'Compound1_ID' in idx:
                add_map(cv('Compound1_ID'), cv('Compound1_name'), 'table_s2_combo')
                add_map(cv('Compound2_ID'), cv('Compound2_name'), 'table_s2_combo')
            if 'Drug_ID' in idx:
                add_map(cv('Drug_ID'), cv('Drug_name'), 'table_s2_single')
    wb.close()
for f in ['gdsc-007_matrix_results.csv', 'gdsc-008_matrix_results.csv', 'gdsc-009_matrix_results.csv', 'gdsc010_matrix_results.csv', 'sandpiper-01_matrix_results.csv']:
    p = os.path.join(TMP, 'extq1', f)
    d = pd.read_csv(p, usecols=['lib1_ID', 'lib1_name', 'lib2_ID', 'lib2_name'], low_memory=False)
    for a, an, b, bn in d.itertuples(index=False):
        add_map(a, an, 'matrix_results'); add_map(b, bn, 'matrix_results')
p = os.path.join(TMP, 'extq1', 'GDSC2_fitted_dose_response.csv')
d = pd.read_csv(p, usecols=['DRUG_ID', 'DRUG_NAME'], low_memory=False)
for a, an in d.itertuples(index=False):
    add_map(a, an, 'gdsc2_single')

# QC barcode -> scan/date
qc = {}
for fn in ['plate_qc_Sandpiper-01_13May22.csv', 'plate_qc_Sandpiper-02_13May22.csv']:
    p = os.path.join(TMP, 'extq1', fn)
    qd = pd.read_csv(p, usecols=['BARCODE', 'SCAN_ID', 'QC_FLAG', 'QC_UPLOAD_DATE'])
    for r in qd.itertuples(index=False):
        qc[as_id(r.BARCODE)] = (as_id(r.SCAN_ID), str(r.QC_FLAG), str(r.QC_UPLOAD_DATE))
qc_dates = sorted({v[2] for v in qc.values()})
qc_scans = {v[0] for v in qc.values()}
qc_flags = Counter(v[1] for v in qc.values())

# raw scan
pair_mats = defaultdict(set)
pair_date_mats = defaultdict(lambda: defaultdict(set))
pair_scan_mats = defaultdict(lambda: defaultdict(set))
pair_barcodes = defaultdict(set)
raw_cells = set(); raw_drugs = set(); raw_barcodes = set(); raw_directed = set()
raw_rows = 0; unmapped = Counter(); rows_by_line = Counter(); barcodes_by_line = defaultdict(set)
line_norms = {norm(x): x for x in ['A375', 'HT-29', 'PC-14', 'U-2-OS', 'SW620', 'C32', 'MHH-ES-1']}
field_nonempty = Counter()
columns = None
target_fields = ['BARCODE', 'SCAN_ID', 'QC_UPLOAD_DATE', 'DRUGSET_ID', 'CELL_ID', 'cmatrix',
                 'lib1_dose', 'lib1_conc', 'lib2_dose', 'lib2_conc', 'INTENSITY', 'inhibition',
                 'HSA', 'HSA_excess', 'Bliss_additivity', 'Bliss_excess', 'treatment',
                 'experiment', 'occasion', 'technical_replicate', 'replicate', 'treatment_date',
                 'passage', 'culture', 'assay', 'duration', 'SCAN_DATE', 'DATE_CREATED']
with zipfile.ZipFile(raw_path) as z:
    for mn in [i.filename for i in z.infolist() if i.filename.lower().endswith('.csv')]:
        with z.open(mn) as f:
            txt = io.TextIOWrapper(f, encoding='utf-8', errors='replace', newline='')
            rdr = csv.reader(txt); hdr = next(rdr)
            if columns is None:
                columns = hdr
            idx = {c: i for i, c in enumerate(hdr)}
            for row in rdr:
                raw_rows += 1
                cell = norm(row[idx['CELL_LINE_NAME']])
                raw_cells.add(cell)
                bc = as_id(row[idx['BARCODE']]); raw_barcodes.add(bc)
                d1id = as_id(row[idx['lib1_drug_id']]); d2id = as_id(row[idx['lib2_drug_id']])
                d1 = norm(id2name.get(d1id, '')); d2 = norm(id2name.get(d2id, ''))
                if not d1: unmapped[d1id] += 1
                if not d2: unmapped[d2id] += 1
                if not d1 or not d2:
                    continue
                raw_drugs.add(d1); raw_drugs.add(d2)
                raw_directed.add((d1, d2))
                pair = tuple(sorted((d1, d2)))
                cm = as_id(row[idx['cmatrix']])
                key = (cell, pair)
                pair_mats[key].add((bc, cm))
                date, scan = qc.get(bc, ('', '', ''))[2], qc.get(bc, ('', '', ''))[0]
                pair_date_mats[key][date].add((bc, cm))
                pair_scan_mats[key][scan].add((bc, cm))
                pair_barcodes[key].add(bc)
                if cell in line_norms:
                    rows_by_line[line_norms[cell]] += 1
                    barcodes_by_line[line_norms[cell]].add(bc)
                for f in ['INTENSITY', 'inhibition', 'HSA', 'HSA_excess', 'Bliss_additivity', 'Bliss_excess']:
                    if f in idx and row[idx[f]].strip() != '':
                        field_nonempty[f] += 1

# counts
def structural_stats(keys):
    mat_counts = Counter(len(pair_mats[k]) for k in keys)
    exact18 = sum(1 for k in keys if len(pair_mats[k]) == 18)
    exact6x3 = sum(1 for k in keys if len(pair_date_mats[k]) == 6 and all(len(v) == 3 for v in pair_date_mats[k].values()))
    strict6x3 = sum(1 for k in keys if len(pair_date_mats[k]) == 6 and all(
        len({b for b, _ in v}) == 3 and len({s for s in [qc.get(b, ('',))[0] for b, _ in v] if s}) == 3
        for v in pair_date_mats[k].values()))
    date_groups = Counter(len(pair_date_mats[k]) for k in keys)
    return {
        'pair_units': len(keys), 'matrix_records': sum(len(pair_mats[k]) for k in keys),
        'units_ge2_matrices': sum(1 for k in keys if len(pair_mats[k]) >= 2),
        'units_exact_18_matrices': exact18,
        'units_six_qc_dates_x3': exact6x3,
        'units_six_qc_dates_x3_strict_distinct_barcode_and_scan': strict6x3,
        'matrix_records_covered_by_six_x3': sum(len(pair_mats[k]) for k in keys if len(pair_date_mats[k]) == 6 and all(len(v) == 3 for v in pair_date_mats[k].values())),
        'matrix_record_count_dist': {str(k): int(v) for k, v in sorted(mat_counts.items())},
        'qc_date_group_count_dist': {str(k): int(v) for k, v in sorted(date_groups.items())},
    }

all_keys = list(pair_mats.keys())
all_stats = structural_stats(all_keys)
line_keys = {name: [k for k in all_keys if k[0] == norm(name)] for name in line_norms.values()}
line_stats = {name: structural_stats(keys) | {'rows': int(rows_by_line[name]), 'barcodes': len(barcodes_by_line[name])}
              for name, keys in line_keys.items()}
# mapping classification for 7 lines/all
def classify(keys):
    exact = 0; six_not3 = 0; lt6 = 0; gt6 = 0; nodate = 0
    for k in keys:
        n = len(pair_date_mats[k])
        if n == 6 and all(len(v) == 3 for v in pair_date_mats[k].values()):
            exact += 1
        elif n == 6:
            six_not3 += 1
        elif n < 6:
            lt6 += 1
        else:
            gt6 += 1
    return {'exact_six_x3': exact, 'six_dates_not_3': six_not3, 'fewer_than_six_dates': lt6, 'more_than_six_dates': gt6,
            'total': len(keys), 'exact_six_x3_fraction': exact / len(keys) if keys else None}
all_class = classify(all_keys)
line_class = {name: classify(keys) for name, keys in line_keys.items()}

# field verification
expected_mapping_fields = ['experiment', 'occasion', 'technical_replicate', 'replicate', 'treatment_date',
                           'passage', 'culture', 'assay', 'duration', 'SCAN_DATE', 'DATE_CREATED']
present_fields = set(c.lower() for c in columns)
missing_mapping_fields = [f for f in expected_mapping_fields if f.lower() not in present_fields]
field_info = {
    'columns': columns,
    'n_rows': raw_rows,
    'n_cells': len(raw_cells), 'n_drugs': len(raw_drugs), 'n_barcodes': len(raw_barcodes),
    'has_mapping_fields': {f: (f.lower() in present_fields) for f in expected_mapping_fields},
    'missing_mapping_fields': missing_mapping_fields,
    'raw_has_scan_id_column': 'SCAN_ID' in present_fields,
    'raw_has_qc_date_column': 'QC_UPLOAD_DATE' in present_fields,
    'qc': {'n_barcodes': len(qc), 'n_scans': len(qc_scans), 'n_dates': len(qc_dates),
           'date_min': qc_dates[0], 'date_max': qc_dates[-1], 'qc_flags': dict(qc_flags)},
    'field_nonempty_counts': dict(field_nonempty),
    'note': 'QC_UPLOAD_DATE/SCAN_ID are in plate QC by BARCODE, not in raw well CSV; QC dates are upload dates, not treatment dates.',
}
# Table S18 header verification
s18 = {}
with zipfile.ZipFile(supp_path) as z:
    wb = openpyxl.load_workbook(z.open('EMS194010-supplement-Table_S18.xlsx'), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = next(rows)
    s18 = {'sheet': wb.sheetnames[0], 'columns': [str(x) for x in hdr],
           'has_mapping_fields': {f: any(f.lower() in str(x).lower() for x in hdr) for f in expected_mapping_fields},
           'has_caspase_parp_western': any(re.search(r'caspase|parp|western|apopto', str(x), re.I) for x in hdr),
           'has_barcode_cmatrix': any('BARCODE' in str(x).upper() for x in hdr) and any('cmatrix' in str(x).lower() for x in hdr)}
    wb.close()

# code usability (read-only)
def code_file(rel):
    p = os.path.join(TMP, 'data2_github', rel)
    return open(p, encoding='utf-8', errors='replace').read() if os.path.exists(p) else ''

pipeline = code_file(r'CancerRxGene_gdscmatrixanalyser\CancerRxGene-gdscmatrixanalyser-00c04b4\R\pipeline.R')
models = code_file(r'CancerRxGene_gdscmatrixexplorer\CancerRxGene-gdscmatrixexplorer-c5c863e\models.py')
nlme = code_file(r'CancerRxGene_gdscIC50\CancerRxGene-gdscIC50-34d8422\R\nlme_fit_prep.R')
eac = code_file(r'eac54_Large-scale-pan-cancer-screening\eac54-Large-scale-pan-cancer-screening-cef258c\Supplementary Figure 1.Rmd')
code_required = ['DATE_CREATED', 'SCAN_ID', 'DRUGSET_ID', 'CELL_ID', 'BARCODE', 'ASSAY', 'DURATION']
code_usability = {
    'gdscmatrixanalyser_make_plate_list_required': ['BARCODE', 'DATE_CREATED', 'SCAN_ID', 'DRUGSET_ID', 'CELL_ID', 'CELL_LINE_NAME', 'MASTER_CELL_ID', 'COSMIC_ID'],
    'public_raw_has': sorted(x for x in code_required if x.lower() in present_fields),
    'public_raw_missing': sorted(x for x in code_required if x.lower() not in present_fields),
    'qc_can_supply': ['SCAN_ID (via BARCODE join)'],
    'cannot_supply': ['DATE_CREATED/treatment date', 'ASSAY', 'DURATION', 'passage/culture/occurrence labels'],
    'gdscmatrixexplorer_all_replicates': 'same drug pair across projects (models.py:70-97/195-224), not six-occasion technical triplicate/independent culture',
    'gdscIC50_requires': ['DATE_CREATED', 'SCAN_DATE', 'ASSAY', 'DURATION'],
    'eac54': 'figure/analysis Rmd only; no mapping/functional numeric',
    'partial_plate_fields_possible': True,
    'occasion_technical_mapping_possible': False,
}

# readiness / qualification, no effect execution
qual = {
    'analysis_executed': False,
    'qualified_for_same_platform_technical_repeat_structure': {
        'status': 'qualified_with_caveats', 'basis': '4,398 units >=2 matrix records; 301 exact-18; 161 6-QC-date x3; 7 lines 763 units',
        'can_do': 'same-platform technical matrix-to-matrix variability/agreement structure; repeat-level descriptive analyses',
        'cannot_claim': 'independent culture; treatment-date occasion; biological replicate; effect inference'},
    'qualified_for_plate_barcode_qc_date_stratified_descriptive': {
        'status': 'qualified_with_caveats', 'basis': 'BARCODE/SCAN_ID/QC_UPLOAD_DATE/DRUGSET_ID/CELL_ID/cmatrix available',
        'can_do': 'plate/barcode/QC-upload-date stratified descriptive tables; not treatment-date strata'},
    'qualified_for_partial_field_mapping': {
        'status': 'partial', 'basis': 'cell/drug/dose/plate/matrix/QC metadata can be assembled; no experiment/occasion/passage/culture',
        'can_do': 'partial key table for cell/drug/dose/plate/cmatrix/QC date/SCAN',
        'still_missing': ['experiment/occasion', 'technical replicate index', 'treatment date', 'passage', 'culture', 'assay/duration']},
    'still_unavailable_independent_culture': {
        'status': 'still_unavailable', 'reason': 'no passage/culture/flask/independent-day fields; technical triplicate != independent culture'},
    'still_unavailable_official_six_occasion_certification': {
        'status': 'still_unavailable', 'reason': 'QC upload dates are not treatment dates; no official per-barcode occasion mapping'},
    'still_unavailable_functional_fig4_5_6_numeric': {
        'status': 'still_unavailable', 'reason': 'raw/public supplement has no caspase/PARP/western numeric; only paper figures/PDF (DATA1/DATA2)'},
    'still_unavailable_cross_background_transfer_real_decision_mechanism': {
        'status': 'still_unavailable', 'reason': 'no independent culture/functional/decision data; not evaluated in this task'},
    'updated_readiness': 'same_platform_structural_only / dev_only; do not change old verdicts',
}

result = {
    'protocol': {'path_rel': 'research/2026-10-28-az-download-verify/PROTOCOL_DATA3.md', 'sha256': sha(PROTOCOL)},    'script': {'path_rel': 'research/2026-10-28-az-download-verify/verify_data3.py', 'sha256': sha(os.path.join(HERE, 'verify_data3.py'))},
    'manifest': {'path_rel': 'research/2026-10-28-az-download-verify/data3_manifest.json',
                 'sha256': sha(MANIFEST), 'summary': man['summary'],
                 'large_reused': [e for e in man['entries'] if e['state'] == 'reused_local'][:20]},
    'field_verification': field_info,
    'table_s18': s18,
    'mapping_reconstruction': {
        'identity_rule': '(cell_norm, unordered canonical drug pair); matrix record=(BARCODE,cmatrix); pair/date/scan/barcode sets',
        'all_export': all_stats,
        'seven_lines': line_stats,
        'all_classification': all_class,
        'seven_line_classification': line_class,
        'reconstruction_ratios': {
            'all_units_exact_six_x3_fraction': all_class['exact_six_x3'] / all_class['total'] if all_class['total'] else None,
            'all_units_ge2_matrices_fraction': all_stats['units_ge2_matrices'] / all_stats['pair_units'] if all_stats['pair_units'] else None,
            'seven_line_units': sum(v['pair_units'] for v in line_stats.values()),
            'seven_line_exact_six_x3_units': sum(v['units_six_qc_dates_x3'] for v in line_stats.values()),
            'seven_line_exact18_units': sum(v['units_exact_18_matrices'] for v in line_stats.values()),
            'seven_line_exact_six_x3_fraction': (sum(v['units_six_qc_dates_x3'] for v in line_stats.values()) / sum(v['pair_units'] for v in line_stats.values())) if sum(v['pair_units'] for v in line_stats.values()) else None,
            'seven_line_exact18_fraction': (sum(v['units_exact_18_matrices'] for v in line_stats.values()) / sum(v['pair_units'] for v in line_stats.values())) if sum(v['pair_units'] for v in line_stats.values()) else None,
            'seven_line_ge2_fraction': (sum(v['units_ge2_matrices'] for v in line_stats.values()) / sum(v['pair_units'] for v in line_stats.values())) if sum(v['pair_units'] for v in line_stats.values()) else None,
            'all_units_6x3_or_better': all_stats['units_six_qc_dates_x3'],
        },
        'caveats': ['QC_UPLOAD_DATE is upload date, not treatment date', 'cmatrix is matrix position, not replicate label',
                    '6x3 reconstruction is QC-date/barcode structure only; not official six-occasion certification',
                    'technical triplicate != independent culture'],
    },
    'code_usability': code_usability,
    'qualification_readiness': qual,
    'notes': [
        'No response values were analysed; only structural/field counts.',
        'Counts are recomputed from raw ZIP with corrected canonical pair identity; compare to X2/X3 (301/161/4,398/68,710/84,055).',
        'QC date groups are upload dates; no independent-culture or official occasion claim.',
    ],
}
with open(OUT, 'w', encoding='utf-8', newline='\n') as f:
    json.dump(q(result), f, indent=1, sort_keys=True, ensure_ascii=False)
    f.write('\n')
print(json.dumps({'all': all_stats, 'seven_lines': line_stats, 'all_class': all_class,
                  'seven_class': line_class, 'field_missing': missing_mapping_fields,
                  's18': s18, 'code': code_usability, 'readiness': qual}, ensure_ascii=False, indent=1)[:12000])
