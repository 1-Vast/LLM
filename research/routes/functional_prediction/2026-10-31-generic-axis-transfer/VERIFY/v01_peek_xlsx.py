"""Print sheet names and the first rows of an xlsx file (layout inspection only)."""
import sys
import openpyxl

path = sys.argv[1]
nrows = int(sys.argv[2]) if len(sys.argv) > 2 else 15
wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
for ws in wb.worksheets:
    print("SHEET", repr(ws.title), "dims:", ws.max_row, "x", ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= nrows:
            break
        cells = list(row)
        # trim trailing None
        while cells and cells[-1] is None:
            cells.pop()
        print(i, len(cells), [c for c in cells[:20]], "..." if len(cells) > 20 else "", cells[-5:] if len(cells) > 20 else "")
wb.close()
