from pathlib import Path
import json
import pymupdf
from PIL import Image
from pptx import Presentation
root=Path(r'D:\MAESTRO\.build-1009\revision')
pdf=pymupdf.open(root/'1009-white-preview.pdf')
for i,page in enumerate(pdf):
    page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5)).save(root/f'white-{i+1:02}.png')
for start in range(0,len(pdf),3):
    canvas=Image.new('RGB',(1440,810*3),'#e5e9e9')
    for j in range(min(3,len(pdf)-start)):
        im=Image.open(root/f'white-{start+j+1:02}.png')
        canvas.paste(im,(0,j*810))
    canvas.save(root/f'review-{start//3+1}.png')
fit=json.loads((root/'fit.json').read_text(encoding='utf-8-sig'))
overflow=[r for r in fit if r['bound']>r['height']+3]
print('Slides:',len(pdf),'Text overflow:',json.dumps(overflow,ensure_ascii=False))
p=Presentation(r'D:\MAESTRO\output-1009\1009.pptx')
bad=[]
for i,s in enumerate(p.slides):
    if 1<=i<=11:
        text=' '.join(sh.text for sh in s.shapes if sh.has_text_frame)+s.notes_slide.notes_text_frame.text
        if 'MAESTRO' in text:bad.append(i+1)
print('MAESTRO references in independent sections:',bad)
print('Native tables:',sum(sh.has_table for s in p.slides for sh in s.shapes))
print('Notes:',[len(s.notes_slide.notes_text_frame.text) for s in p.slides])
